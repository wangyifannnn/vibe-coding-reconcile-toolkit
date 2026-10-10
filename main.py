from __future__ import annotations

import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_FILE = OUTPUT_DIR / "reconciliation_result.xlsx"


def clean_phone(value: str) -> tuple[str, bool]:
    raw = value.strip()
    if raw.startswith("+86"):
        raw = raw[3:]
    phone = re.sub(r"[\s-]", "", raw)
    return phone, bool(re.fullmatch(r"\d{11}", phone))


def clean_date(value: str) -> str:
    normalized = value.strip().replace("/", "-").replace(".", "-")
    if not normalized:
        return ""
    try:
        return datetime.fromisoformat(normalized).strftime("%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"无法识别下单时间：{value}") from exc


def parse_amount(value: str) -> Decimal | None:
    text = value.strip()
    if not text:
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"无法识别金额：{value}") from exc


def locate_input(required_columns: set[str]) -> Path:
    matches: list[Path] = []
    for path in INPUT_DIR.glob("*.csv"):
        columns = set(pd.read_csv(path, dtype=str, keep_default_na=False, nrows=0).columns)
        if required_columns.issubset(columns):
            matches.append(path)
    if len(matches) != 1:
        names = "、".join(sorted(required_columns))
        raise FileNotFoundError(f"input 目录中应有且仅有一个包含以下字段的 CSV：{names}")
    return matches[0]


def load_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")


def build_result() -> tuple[pd.DataFrame, pd.DataFrame]:
    order_columns = {"订单号", "客户手机", "支付金额", "下单时间", "收货地址"}
    refund_columns = {"原订单号", "退款金额"}
    orders = load_csv(locate_input(order_columns))
    refunds = load_csv(locate_input(refund_columns))

    # 同一订单号只保留首次出现的记录，再按清洗规则判断是否隔离。
    orders["_订单号"] = orders["订单号"].str.strip()
    orders = orders.drop_duplicates(subset=["_订单号"], keep="first")

    refund_totals: dict[str, Decimal] = {}
    for _, row in refunds.iterrows():
        order_id = row["原订单号"].strip()
        amount = parse_amount(row["退款金额"])
        refund_totals[order_id] = refund_totals.get(order_id, Decimal("0")) + (amount or Decimal("0"))

    cleaned_rows: list[dict[str, object]] = []
    anomaly_rows: list[dict[str, str]] = []
    source_columns = ["订单号", "客户手机", "支付金额", "下单时间", "收货地址"]

    for _, row in orders.iterrows():
        phone, phone_ok = clean_phone(row["客户手机"])
        payment = parse_amount(row["支付金额"])
        reasons: list[str] = []
        if payment is None:
            reasons.append("金额缺失")
        if not row["收货地址"].strip():
            reasons.append("地址缺失")
        if not phone_ok:
            reasons.append("手机号异常")

        if reasons:
            original = {column: row[column] for column in source_columns}
            original["异常原因"] = "；".join(reasons)
            anomaly_rows.append(original)
            continue

        order_id = row["_订单号"]
        refund = refund_totals.get(order_id, Decimal("0"))
        cleaned_rows.append(
            {
                "订单号": order_id,
                "规范手机": phone,
                "规范下单日期": clean_date(row["下单时间"]),
                "原始支付金额": float(payment),
                "退款金额": float(refund),
                "实际营收": float(payment - refund),
            }
        )

    cleaned_columns = ["订单号", "规范手机", "规范下单日期", "原始支付金额", "退款金额", "实际营收"]
    anomaly_columns = source_columns + ["异常原因"]
    return (
        pd.DataFrame(cleaned_rows, columns=cleaned_columns),
        pd.DataFrame(anomaly_rows, columns=anomaly_columns),
    )


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    widths = {
        "Cleaned_Orders": [16, 16, 18, 16, 14, 14],
        "Anomaly_Log": [16, 20, 16, 22, 28, 24],
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for index, width in enumerate(widths[sheet.title], start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        if sheet.title == "Cleaned_Orders":
            for row in sheet.iter_rows(min_row=2, min_col=4, max_col=6):
                for cell in row:
                    cell.number_format = "0.00"
    workbook.save(path)


def main() -> int:
    try:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        cleaned, anomalies = build_result()
        with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
            cleaned.to_excel(writer, sheet_name="Cleaned_Orders", index=False)
            anomalies.to_excel(writer, sheet_name="Anomaly_Log", index=False)
        format_workbook(OUTPUT_FILE)
    except Exception as exc:
        print(f"处理失败：{exc}", file=sys.stderr)
        return 1

    print(f"处理完成：{OUTPUT_FILE}")
    print(f"有效订单：{len(cleaned)} 条；异常订单：{len(anomalies)} 条")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
